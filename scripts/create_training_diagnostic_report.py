from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "docs" / "训练效果差检测报告.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_col_widths(table, widths):
    widths_twips = [int(round(width * 1440)) for width in widths]
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_twips)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "110")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(round(width * 1440))))
            tc_w.set(qn("w:type"), "dxa")


def style_run(run, size=10.5, bold=False, color="000000", italic=False):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def format_paragraph(paragraph, before=0, after=6, line=1.25, align=None):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_text(doc, text, style=None, before=0, after=6, line=1.25, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    format_paragraph(p, before=before, after=after, line=line)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        style_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        style_run(r2)
    else:
        r = p.add_run(text)
        style_run(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.65 + level * 0.45)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    format_paragraph(p, after=3, line=1.18)
    style_run(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    format_paragraph(p, before=10 if level == 1 else 6, after=5, line=1.1)
    r = p.add_run(text)
    style_run(r, size=15 if level == 1 else 12, bold=True)
    return p


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_col_widths(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        format_paragraph(p, after=0, line=1.05, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(str(text))
        style_run(r, size=font_size, bold=True, color="FFFFFF")
    for ridx, rowdata in enumerate(rows):
        row = table.add_row()
        for i, text in enumerate(rowdata):
            cell = row.cells[i]
            set_cell_shading(cell, "F2F6FA" if ridx % 2 else "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.LEFT if i == 0 or len(str(text)) > 28 else WD_ALIGN_PARAGRAPH.CENTER
            format_paragraph(p, after=0, line=1.08, align=align)
            r = p.add_run(str(text))
            style_run(r, size=font_size)
    set_col_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    style_run(run, size=8.5, color="666666")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run2 = paragraph.add_run(" 页")
    style_run(run2, size=8.5, color="666666")


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        styles[name].font.name = "Microsoft YaHei"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(0, 0, 0)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(header.add_run("YOLO Fire 训练效果检测报告"), size=8.5, color="666666")
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(title, before=22, after=8, line=1.1)
    r = title.add_run("火灾烟雾检测训练效果差检测报告")
    style_run(r, size=22, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(sub, after=14, line=1.15)
    style_run(sub.add_run("基于远程训练仓库和 D Fire 数据集的诊断结论"), size=11, color="555555")

    add_table(doc, ["项目", "内容"], [
        ("报告日期", "2026 年 9 月 21 日"),
        ("检查对象", "YOLO11n baseline、P2、SRDG、DCBR、ACR 训练结果"),
        ("远程仓库", "/home/cmj/Documents/FireSmoke/Fire-YOLOv11/"),
        ("远程代码版本", "ce4a405 arc_seed0"),
        ("数据集", "D Fire，train 15485，val 1722，test 4294"),
        ("结论级别", "诊断报告，不替代最终 test 或跨域结论"),
    ], [1.3, 5.45], font_size=9.4)

    add_heading(doc, "执行摘要", 1)
    add_text(doc, "当前训练效果差，数据集问题是重要原因，但不是唯一原因。远程误差分析显示，主要瓶颈不是 fire 和 smoke 互相分类错误，而是小目标定位、整体框质量、重复框、置信度排序和背景误报。数据集同时存在高比例空图、跨 split 近重复候选、目标尺度过小、越界框和标注粒度可能不一致等风险。继续堆叠 ACR、DCBR 等结构模块，暂时不能解决这些基础问题。")
    add_text(doc, "当前最优策略是暂停新模块扩展，先固定数据版本并完成分组划分、近重复审查和标注质量审查；然后在清洗后的数据上重新训练严格可比的 baseline 和 P2，再决定是否继续模块研究。")

    add_heading(doc, "诊断范围与证据", 1)
    add_text(doc, "本报告基于远程训练机上的原始数据、训练日志、权重和验证集误差分析产物整理。训练结果均来自 D Fire 验证集，模型选择使用验证 mAP50-95；COCO 尺寸分层、固定工作点误差和低误报召回属于统一验证集上的诊断，不应直接当作测试集或部署保证。")
    add_bullet(doc, "训练配置：100 epochs、640 输入、batch 32、SGD、AMP=false，P2、SRDG、DCBR、ACR 使用同协议。")
    add_bullet(doc, "远程机器：RTX 3090，Python 3.10.12，Ultralytics 8.4.42，PyTorch 2.11.0。")
    add_bullet(doc, "分析报告：/home/cmj/Documents/FireSmoke/Fire-YOLOv11/artifacts/analysis/review-20260920-ce4a405/REPORT.md。")
    add_bullet(doc, "本报告没有使用测试集结果决定模型或超参数。")

    add_heading(doc, "训练结果概览", 1)
    add_table(doc, ["模型", "seed", "最佳轮次", "mAP50", "mAP50-95", "P", "R", "初步判断"], [
        ("P2", "0", "100", "80.345%", "48.384%", "78.426%", "74.535%", "当前单次最高"),
        ("P2", "1", "94", "80.206%", "47.708%", "78.791%", "73.400%", "低于 seed0"),
        ("P2", "2", "95", "79.358%", "47.609%", "79.753%", "71.573%", "召回下降"),
        ("P2 SRDG", "0", "94", "79.863%", "48.098%", "82.104%", "71.595%", "P 上升但 R 下降"),
        ("P2 DCBR", "0", "78", "80.250%", "48.160%", "78.684%", "73.126%", "低于 P2 seed0"),
        ("P2 DCBR", "2", "93", "79.811%", "48.077%", "77.418%", "74.370%", "高于 P2 seed2"),
        ("P2 ACR", "0", "74", "79.580%", "47.854%", "80.426%", "72.223%", "低于 P2 0.530 个百分点"),
    ], [1.0, 0.42, 0.62, 0.62, 0.78, 0.58, 0.58, 1.8], font_size=8.4)
    add_text(doc, "P2 三个 seed 的 mAP50-95 为 47.900% ± 0.422%。DCBR 只有两个 seed，配对平均变化约为 +0.122 个百分点且方向相反，尚不能证明稳定收益。ACR seed0 相比同 seed P2 下降 0.530 个百分点，超过预先设置的 0.300 个百分点筛查阈值。")

    add_heading(doc, "数据集诊断", 1)
    add_heading(doc, "空图比例较高", 2)
    add_text(doc, "D Fire 中约 45%–47% 图片没有目标框。空图有助于抑制误报，但会把训练目标推向背景保守策略，尤其会伤害低对比度烟雾和远距离小火焰的召回。")
    add_table(doc, ["Split", "图像数", "空标签图", "空图比例", "处理建议"], [
        ("train", "15485", "7043", "45.5%", "保留，但按来源和场景审查"),
        ("val", "1722", "790", "45.9%", "固定版本并统计背景报警率"),
        ("test", "4294", "2005", "46.7%", "冻结，不用于反复调参"),
    ], [0.8, 1.0, 1.1, 1.0, 2.85], font_size=9.0)

    add_heading(doc, "目标尺度不适合只用 640 输入", 2)
    add_text(doc, "严格按原图框面积统计，火焰目标中约 31% 属于 COCO small；缩放到 640 后，约 23%–27% 的火焰目标短边小于 16 像素。验证集上 P2 的火焰 AP-small 只有 23.656%，而火焰 AP-large 为 57.528%，说明总体指标被小目标和定位误差显著拖低。")
    add_table(doc, ["类别", "训练框数", "small 比例", "640 后短边小于 16 像素", "含义"], [
        ("火焰", "10677", "约 30.6%", "约 25.5%", "小目标和框定位是主要压力"),
        ("烟雾", "8582", "约 3.7%", "约 4.4%", "更多是弥散边界和低误报召回问题"),
    ], [0.9, 1.0, 1.2, 1.8, 2.0], font_size=8.8)

    add_heading(doc, "数据来源之间存在明显差异", 2)
    add_text(doc, "训练数据由 AoF、WEB 和 PublicDataset 等来源组成。不同来源的空图比例和目标密度不同，随机切分会让模型同时学习来源风格，导致结果对场景和 seed 敏感。")
    add_table(doc, ["来源", "训练图像", "空图", "有目标图", "空图比例"], [
        ("AoF", "6031", "3604", "2427", "59.8%"),
        ("WEB", "8490", "3318", "5172", "39.1%"),
        ("PublicDataset", "964", "121", "843", "12.6%"),
    ], [1.55, 1.25, 1.15, 1.35, 1.45], font_size=9.0)

    add_heading(doc, "标注质量需要复核", 2)
    add_text(doc, "数据准备记录显示，train 有 14 个、test 有 12 个图像因非法框被排除。对原始标签做严格扫描时，还发现 train、val、test 分别约有 292、25、62 个框越过图像边界。越界框、零面积框以及整体烟雾框与局部烟雾框粒度不一致，都会直接影响 mAP50-95 和 AP75。")
    add_text(doc, "这些问题不代表所有异常框都一定是错误标注，但必须建立统一的裁剪、重标或排除规则，并重新生成版本化的数据清单。")

    add_heading(doc, "跨 split 近重复风险", 2)
    add_text(doc, "按 AoF、WEB 和 PublicDataset 的相邻编号检查，发现 3057 对相邻编号图片跨越了 train、val 或 test。进一步检查得到 743 对 dHash 距离不超过 2，其中 101 对解码后的平均像素差小于 1，585 对平均像素差小于 5；没有发现字节级完全相同的图片。")
    add_text(doc, "这不是重复泄漏的最终证明，但已经足以作为强近重复候选。相邻视频帧或同一场景图片跨 split 会让验证结果偏乐观，也会使随机 split 的 seed 比较失去场景独立性。当前 metadata 还记录 image_content_hashed=false，说明原流程没有完成严格内容级去重。")

    add_heading(doc, "模型错误诊断", 1)
    add_heading(doc, "主要错误来自定位和背景而不是类别混淆", 2)
    add_table(doc, ["模型", "TP", "FN", "定位 FP", "重复 FP", "低重叠背景 FP", "空图报警"], [
        ("P2 seed0", "1588", "503", "280", "86", "186", "8 / 790"),
        ("DCBR seed0", "1590", "501", "311", "99", "190", "12 / 790"),
        ("ACR seed0", "1581", "510", "312", "101", "201", "15 / 790"),
    ], [1.0, 0.62, 0.62, 0.9, 0.8, 1.55, 1.1], font_size=8.8)
    add_text(doc, "类别混淆误报数量很少，而定位误报、重复框和低重叠背景误报更突出。DCBR 和 ACR 在该固定验证工作点上没有降低背景错误，反而比 P2 产生更多空图报警。")

    add_heading(doc, "低误报工作点下烟雾召回下降", 2)
    add_text(doc, "在每类负图经验 FPR 不超过 1% 的验证集工作点下，P2 的 smoke 图像召回为 70.225%，DCBR 为 62.752%，ACR 为 66.667%。这说明模块带来的精度上升不能直接转化为低误报条件下的烟雾预警能力。阈值校准只能改变工作点，不能凭空提高完整检测 AP。")
    add_table(doc, ["模型", "fire 图像召回", "smoke 图像召回", "判断"], [
        ("P2", "95.662%", "70.225%", "当前较均衡"),
        ("DCBR", "96.095%", "62.752%", "烟雾召回明显下降"),
        ("ACR", "93.926%", "66.667%", "两类均未优于 P2"),
    ], [1.25, 1.4, 1.45, 2.65], font_size=9.0)

    add_heading(doc, "模块没有解决主要失败模式", 2)
    add_text(doc, "DCBR seed0 在小火焰 AP 上有约 1.226 个百分点的局部收益，但中大火焰和严格定位指标下降，因此不是整体突破。ACR 的已训练权重中，上下文细节残差分支对推理 AP 的影响约为 0.001 个百分点，说明该分支在当前训练中几乎没有产生有效贡献。")
    add_text(doc, "因此当前不应继续通过增加模块数量来掩盖数据和任务定义问题。应先解释整体框、局部框、重复框和置信度排序错误。")

    add_heading(doc, "根因判断", 1)
    add_table(doc, ["问题", "证据强度", "对结果的影响", "当前判断"], [
        ("小火焰尺寸过小", "强", "AP-small 和 AP75 显著低", "主要根因之一"),
        ("空图比例高", "强", "背景抑制与烟雾召回冲突", "重要训练张力"),
        ("跨 split 近重复", "中到强", "验证独立性不足", "必须修正协议"),
        ("越界或异常标注", "中", "损害定位指标", "需要版本化清洗"),
        ("来源风格差异", "中到强", "seed 和场景敏感", "需要分组划分"),
        ("模型结构不够复杂", "弱到中", "当前证据不支持", "暂不优先"),
        ("类别映射完全错误", "弱", "当前证据不支持", "暂未发现"),
    ], [1.55, 1.0, 2.0, 2.2], font_size=8.7)

    add_heading(doc, "下一步工作计划", 1)
    add_heading(doc, "第一阶段 数据和划分审计", 2)
    add_bullet(doc, "保留当前所有权重和结果，不覆盖现有实验目录。")
    add_bullet(doc, "对 743 对强近重复候选进行人工确认，并建立场景、视频或来源组 ID。")
    add_bullet(doc, "按场景或视频分组重新划分 train、val、test，禁止同一序列跨 split。")
    add_bullet(doc, "逐一审查越界框、零面积框、局部框和整体框，明确统一标注规则。")
    add_bullet(doc, "固定新的 manifest、排除清单和 SHA256，形成新的数据版本。")

    add_heading(doc, "第二阶段 建立可归因的基线", 2)
    add_bullet(doc, "在清洗后的数据上用相同 seed、初始化、优化器、AMP 和训练预算重训 baseline 与 P2。")
    add_bullet(doc, "先比较 640 与 960 输入的单因素变化，不同时改变模块、loss、epoch 和增强。")
    add_bullet(doc, "报告 AP50、AP75、AP-small、AP-medium、AP-large、低 FPR smoke recall 和 FP/image。")
    add_bullet(doc, "只有当 P2 在分组测试上稳定优于 baseline，才继续讨论 P2 后续模块。")

    add_heading(doc, "第三阶段 决定是否继续模块", 2)
    add_bullet(doc, "ACR 暂停 seed1、seed2 和完整消融，先检查上下文分支梯度、激活尺度和输出投影更新量。")
    add_bullet(doc, "DCBR 保留为候选，补齐 seed1 后再判断是否存在稳定收益，不能只保留有利 seed。")
    add_bullet(doc, "若错误仍以整体框和重复框为主，优先研究标注定义、框合并和训练目标，不再堆叠注意力模块。")

    add_heading(doc, "验收标准", 2)
    add_table(doc, ["项目", "最低要求", "说明"], [
        ("数据清单", "train、val、test 无同场景跨 split", "至少完成候选近重复人工确认"),
        ("标签质量", "非法框有统一处理规则", "保留排除清单和原因"),
        ("可归因性", "baseline 与候选只改变一个因素", "统一初始化和训练预算"),
        ("检测质量", "同时报告 AP75 和 AP-small", "不能只看 mAP50"),
        ("低误报能力", "报告每类 1% FPR 下召回", "重点关注 smoke recall"),
        ("部署价值", "报告真实设备 latency、p95 和显存", "训练时间不能替代推理效率"),
    ], [1.45, 2.25, 3.05], font_size=8.8)

    add_heading(doc, "最终结论", 1)
    add_text(doc, "当前结果不足以支持继续无目的地增加网络模块。数据集中的小目标、空图、来源差异、近重复风险和标注异常，已经足以影响训练稳定性和评价可信度；模型误差又主要集中在定位、重复和背景误报，恰好不是 ACR 或 DCBR 当前设计直接解决的问题。")
    add_text(doc, "下一步应先完成数据版本和分组评估协议，再用严格可比的 baseline 与 P2 重新建立基准。如果清洗后的结果仍然显示小火焰和整体框问题，再单独设计针对输入分辨率、采样策略或标注定义的实验。")

    add_heading(doc, "参考产物", 1)
    add_bullet(doc, "远程诊断报告：/home/cmj/Documents/FireSmoke/Fire-YOLOv11/artifacts/analysis/review-20260920-ce4a405/REPORT.md")
    add_bullet(doc, "数据元数据：/home/cmj/Documents/FireSmoke/Fire-YOLOv11/artifacts/data/dfire/metadata.json")
    add_bullet(doc, "训练结果目录：/home/cmj/Documents/FireSmoke/Fire-YOLOv11/experiments/yolo11n/")
    add_bullet(doc, "分析结果：summary.json、descriptive-errors.json、acr-sensitivity.json、module-probes.json")

    doc.core_properties.title = "火灾烟雾检测训练效果差检测报告"
    doc.core_properties.subject = "D Fire 数据集和 YOLO11n 训练结果诊断"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = "基于远程训练仓库的只读诊断结果"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
